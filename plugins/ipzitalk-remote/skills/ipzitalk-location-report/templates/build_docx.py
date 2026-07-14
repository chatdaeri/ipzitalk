#!/usr/bin/env python3
"""result.json -> 세부 입지 보고서 DOCX.

사용:  python3 build_docx.py result.json out.docx [map.png]
      지도는 map.url 페이지를 헤드리스 Chrome으로 캡처해 자동 삽입한다(mapshot.py).
필요:  pip install python-docx

레이아웃은 고정이다. 바뀌는 건 result.json 뿐이다.
값이 없으면 "정보없음". 절대 0이나 추정값으로 채우지 않는다.
"""
import json
import sys

from mapshot import resolve_map_image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

POINT = RGBColor(0xCE, 0x17, 0x3D)
INK = RGBColor(0x2B, 0x2D, 0x3A)
INK_SUB = RGBColor(0x6B, 0x72, 0x80)
FONT = "S-Core Dream"  # 미설치 환경에서는 시스템 기본 고딕으로 폴백된다
NA = "정보없음"


def style(run, size=10, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def para(doc, text, size=10, bold=False, color=INK, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    style(p.add_run(text), size, bold, color)
    return p


def cell_text(v):
    if v is None:
        return NA, True
    if isinstance(v, dict):
        return str(v.get("n", NA)), True
    return str(v), False


def stars(n):
    if n is None:
        return "판정 유보"
    return "★" * int(n) + "☆" * (3 - int(n))


def add_table(doc, columns, rows):
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = "Table Grid"
    for c, name in enumerate(columns):
        cell = t.rows[0].cells[c]
        cell.text = ""
        style(cell.paragraphs[0].add_run(str(name)), 9, True)
    for row in rows:
        cells = t.add_row().cells
        for c in range(len(columns)):
            text, numeric = cell_text(row[c] if c < len(row) else None)
            cells[c].text = ""
            p = cells[c].paragraphs[0]
            if numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            style(p.add_run(text), 9, color=INK_SUB if text == NA else INK)


def build(d, out_path, map_png=None):
    map_png = resolve_map_image(d, map_png, prefix="ipzitalk-map-docx")
    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10)

    para(doc, d.get("brand") or "IPZI TALK", 9, True, POINT, space_after=2)
    para(doc, d.get("title") or "세부 입지 보고서", 22, True, INK, space_after=2)
    para(doc, d.get("subtitle") or "", 10, color=INK_SUB, space_after=8)
    pills = " · ".join(p.get("label", "") for p in (d.get("pills") or []))
    if pills:
        para(doc, pills, 10, True, POINT)

    m = d.get("map") or {}
    if map_png:
        doc.add_picture(map_png, width=Inches(6.0))
    if m.get("url"):
        para(doc, "지도(동적): " + m["url"], 9, color=INK_SUB, space_after=2)
        for line in [m.get("caption"), m.get("ttlNote")]:
            if line:
                para(doc, line, 8, color=INK_SUB, space_after=2)

    para(doc, "종합 요약", 14, True, INK, space_after=4)
    para(doc, d.get("summary") or NA)
    axes = d.get("axes") or []
    if axes:
        add_table(doc, ["축", "등급", "근거"],
                  [[a.get("label") or "", stars(a.get("stars")), a.get("note") or None] for a in axes])
        doc.add_paragraph()

    for sec in d.get("sections") or []:
        para(doc, sec.get("title") or "", 14, True, INK, space_after=2)
        head = " · ".join(x for x in [sec.get("gradeLabel"), sec.get("band")] if x)
        if head:
            para(doc, head, 9, color=INK_SUB)
        add_table(doc, sec.get("columns") or [], sec.get("rows") or [])
        # 누락·정합성 단서는 경고가 아니라 각주다. 회색으로 작게.
        foot = [x for x in [sec.get("note")] if x]
        if sec.get("truncated"):
            foot.append("목록 불완전 — 검색 건수 제한(45건 캡)으로 일부 누락 가능")
        if foot:
            para(doc, " · ".join(foot), 7, color=INK_SUB)
        doc.add_paragraph()

    wide = d.get("wide")
    if wide:
        para(doc, wide.get("title") or "광역 축", 14, True, INK, space_after=2)
        if wide.get("note"):
            para(doc, wide["note"], 9, color=INK_SUB)
        add_table(doc, wide.get("columns") or [], wide.get("rows") or [])
        doc.add_paragraph()

    if d.get("manual"):
        para(doc, "수기 편집 영역", 12, True, POINT, space_after=2)
        for x in d["manual"]:
            para(doc, "· " + x, 9, color=INK_SUB, space_after=2)
        doc.add_paragraph()

    para(doc, "유의사항 · 출처", 12, True, INK, space_after=2)
    for x in d.get("hedges") or []:
        para(doc, "· " + x, 9, space_after=2)
    foot = []
    if d.get("collectedAt"):
        foot.append("조회일 " + d["collectedAt"])
    if d.get("sources"):
        foot.append("출처: " + " · ".join(d["sources"]))
    if foot:
        para(doc, " · ".join(foot), 8, color=INK_SUB)
    if d.get("disclaimer"):
        para(doc, d["disclaimer"], 8, color=INK_SUB)

    doc.save(out_path)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "result.json"
    dst = sys.argv[2] if len(sys.argv) > 2 else "result.docx"
    png = sys.argv[3] if len(sys.argv) > 3 else None
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    build(data, dst, png)
    print(f"saved: {dst}")
