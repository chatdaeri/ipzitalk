#!/usr/bin/env python3
"""result.json -> 세부 입지 보고서 PPTX (발표용 · HTML 리포트와 같은 다크 스타일).

사용:  python3 build_pptx.py result.json out.pptx [map.png]
필요:  pip install python-pptx

디자인 토큰은 templates/result.html 의 다크 테마와 1:1로 맞춘다.
레이아웃은 고정이다. 바뀌는 건 result.json 뿐이다.
값이 없으면 "정보없음". 절대 0이나 추정값으로 채우지 않는다.
"""
import json
import sys

from mapshot import resolve_map_image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

# ── result.html 다크 테마 토큰 ──
BG = RGBColor(0x0C, 0x0F, 0x14)
CARD = RGBColor(0x16, 0x1B, 0x22)
CHIP = RGBColor(0x1E, 0x24, 0x2E)
ZEBRA = RGBColor(0x18, 0x1D, 0x25)
LINE = RGBColor(0x2A, 0x31, 0x3D)
INK = RGBColor(0xE8, 0xEC, 0xF2)
SUB = RGBColor(0x9A, 0xA4, 0xB2)
SUB2 = RGBColor(0x6B, 0x76, 0x84)
BRAND = RGBColor(0x5B, 0x8B, 0xFF)
TONES = {  # tone → (별·등급 글자색, 등급 pill 배경)
    "g": (RGBColor(0x38, 0xCF, 0x7A), RGBColor(0x0F, 0x2A, 0x1D)),
    "y": (RGBColor(0xE8, 0xB5, 0x3F), RGBColor(0x2B, 0x24, 0x10)),
    "o": (RGBColor(0xF0, 0x81, 0x3E), RGBColor(0x2C, 0x1D, 0x12)),
    "r": (RGBColor(0xF0, 0x5A, 0x5F), RGBColor(0x30, 0x16, 0x19)),
    "x": (RGBColor(0x8B, 0x95, 0xA3), RGBColor(0x21, 0x28, 0x32)),
}
FONT = "Apple SD Gothic Neo"  # 미설치 환경에서는 시스템 기본 고딕으로 폴백된다
MONO = "Menlo"

W, H = Pt(720), Pt(405)  # 16:9
M = Pt(36)
NA = "정보없음"
ROWS_PER_SLIDE = 8  # 발표용: 한 장에 너무 많이 담지 않는다


def tone_of(t):
    return TONES.get(t or "x", TONES["x"])


def txt(s, x, y, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, font=FONT):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), bold, color, font
    return box


def card(s, x, y, w, h, fill=CARD, line=LINE, radius=0.04):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    return shp


def label_width(label, size=11):
    """한글은 전각이라 반각의 두 배로 잡는다. 좁게 잡으면 칩 안에서 줄바꿈된다."""
    units = sum(1.0 if ord(ch) > 0x2000 else 0.55 for ch in label)
    return Pt(22 + units * size)


def chip_text(s, x, y, label, fg=INK, bg=CHIP, size=11):
    w = label_width(label, size)
    shp = card(s, x, y, w, Pt(24), fill=bg, line=None, radius=0.5)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), True, fg, FONT
    return w


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def heading(s, title, grade=None, tone=None):
    txt(s, M, Pt(26), Pt(480), Pt(28), title, size=21, bold=True)
    if grade:
        fg, bgc = tone_of(tone)
        chip_text(s, W - M - label_width(grade), Pt(28), grade, fg=fg, bg=bgc, size=11)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, Pt(58), Pt(26), Pt(2.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND
    bar.line.fill.background()
    bar.shadow.inherit = False


def cell_text(v):
    """rows 셀: null=정보없음, {"n": "279m"}=숫자셀(우측정렬·모노)."""
    if v is None:
        return NA, True
    if isinstance(v, dict):
        return str(v.get("n", NA)), True
    return str(v), False


def stars(n):
    return "판정 유보" if n is None else "★" * int(n) + "☆" * (3 - int(n))


def cell_border(cell, color=LINE):
    """python-pptx는 셀 테두리 API가 없다. HTML처럼 아래쪽 헤어라인만 남기고 나머지는 지운다."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for tag in ("a:lnL", "a:lnR", "a:lnT"):
        ln = tcPr.makeelement(qn(tag), {"w": "0"})
        ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.append(ln)
    ln = tcPr.makeelement(qn("a:lnB"), {"w": "9525", "cap": "flat"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": f"{color}"})
    fill.append(clr)
    ln.append(fill)
    tcPr.append(ln)


def add_table(s, x, y, w, columns, rows):
    tbl = s.shapes.add_table(len(rows) + 1, len(columns), x, y, w, Pt(21) * (len(rows) + 1)).table
    tbl.first_row = False
    for c, name in enumerate(columns):
        cell = tbl.cell(0, c)
        cell.text = str(name)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CHIP
        r = cell.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(10), True, SUB, FONT
        cell_border(cell)
    for i, row in enumerate(rows, start=1):
        for c in range(len(columns)):
            text, numeric = cell_text(row[c] if c < len(row) else None)
            cell = tbl.cell(i, c)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD if i % 2 else ZEBRA
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
            r = p.runs[0]
            r.font.size = Pt(10)
            r.font.name = MONO if numeric else FONT
            r.font.color.rgb = SUB2 if text == NA else INK
            cell_border(cell)
    return tbl


def slide_cover(prs, d):
    s = blank(prs)
    card(s, M, Pt(56), W - M * 2, Pt(256))
    txt(s, M + Pt(26), Pt(86), Pt(400), Pt(16), d.get("brand") or "IPZI TALK", size=11, bold=True, color=BRAND, font=MONO)
    txt(s, M + Pt(26), Pt(112), W - M * 2 - Pt(52), Pt(46), d.get("title") or "세부 입지 보고서", size=30, bold=True)
    txt(s, M + Pt(26), Pt(164), W - M * 2 - Pt(52), Pt(22), d.get("subtitle") or "", size=13, color=SUB)

    x = M + Pt(26)
    for p in d.get("pills") or []:
        x += chip_text(s, x, Pt(208), p.get("label") or "") + Pt(8)

    foot = [f for f in [d.get("collectedAt") and "조회일 " + d["collectedAt"], " · ".join(d.get("sources") or [])] if f]
    txt(s, M + Pt(26), Pt(252), W - M * 2 - Pt(52), Pt(16), " · ".join(foot), size=9, color=SUB2)
    m = d.get("map") or {}
    if m.get("url"):
        txt(s, M, Pt(330), W - M * 2, Pt(30), "지도(동적) " + m["url"] + "\n" + (m.get("ttlNote") or ""), size=7, color=SUB2)


def slide_summary(prs, d, map_png):
    """좌(지도) | 우(요약 + 축 카드 3개). 지도가 없으면 우측이 전폭."""
    s = blank(prs)
    heading(s, "종합 요약")

    left_w = Pt(300)
    if map_png:
        pic = s.shapes.add_picture(map_png, M, Pt(76), width=left_w)  # 높이는 원본 비율
        cap = (d.get("map") or {}).get("caption") or ""
        txt(s, M, Pt(76) + pic.height + Pt(6), left_w, Pt(36), cap, size=7, color=SUB2)
        x0 = M + left_w + Pt(20)
    else:
        x0 = M
    right_w = W - M - x0

    box = card(s, x0, Pt(76), right_w, Pt(88), fill=CHIP)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Pt(12)
    tf.margin_top = Pt(10)
    r = tf.paragraphs[0].add_run()
    r.text = d.get("summary") or NA
    r.font.size, r.font.color.rgb, r.font.name = Pt(9.5), INK, FONT

    y = Pt(176)
    for a in d.get("axes") or []:
        fg, _ = tone_of(a.get("tone"))
        card(s, x0, y, right_w, Pt(50))
        txt(s, x0 + Pt(14), y + Pt(7), Pt(60), Pt(16), a.get("label") or "", size=11, bold=True)
        txt(s, x0 + Pt(14), y + Pt(25), Pt(60), Pt(16), stars(a.get("stars")), size=12, bold=True, color=fg)
        txt(s, x0 + Pt(84), y + Pt(9), right_w - Pt(98), Pt(36), a.get("note") or "", size=9, color=SUB)
        y += Pt(56)


def slide_section(prs, sec, page, pages, rows):
    s = blank(prs)
    title = sec.get("title") or ""
    if pages > 1:
        title = f"{title} ({page}/{pages})"
    heading(s, title, sec.get("gradeLabel"), sec.get("tone"))
    if sec.get("band"):
        txt(s, M, Pt(66), W - M * 2, Pt(24), "밴드: " + sec["band"], size=7.5, color=SUB2)
    card(s, M - Pt(6), Pt(94), W - M * 2 + Pt(12), Pt(21) * (len(rows) + 1) + Pt(12))
    add_table(s, M, Pt(100), W - M * 2, sec.get("columns") or [], rows)

    # 누락·정합성 단서는 경고가 아니라 각주다. 회색으로 작게, 구석에.
    foot = [f for f in [sec.get("note")] if f]
    if sec.get("truncated"):
        foot.append("목록 불완전 — 검색 건수 제한(45건 캡)으로 일부 누락 가능")
    if foot:
        txt(s, M, H - Pt(36), W - M * 2, Pt(26), " · ".join(foot), size=7, color=SUB2)


def slide_wide(prs, wide):
    s = blank(prs)
    heading(s, wide.get("title") or "광역 축")
    if wide.get("note"):
        txt(s, M, Pt(66), W - M * 2, Pt(24), wide["note"], size=7.5, color=SUB2)
    rows = wide.get("rows") or []
    card(s, M - Pt(6), Pt(94), W - M * 2 + Pt(12), Pt(21) * (len(rows) + 1) + Pt(12))
    add_table(s, M, Pt(100), W - M * 2, wide.get("columns") or [], rows)


def slide_notes(prs, d):
    s = blank(prs)
    heading(s, "유의사항 · 출처")
    hedges = d.get("hedges") or []
    manual = d.get("manual") or []

    card(s, M, Pt(74), W - M * 2, Pt(16) * len(hedges) + Pt(20))
    y = Pt(84)
    for h in hedges:
        txt(s, M + Pt(14), y, W - M * 2 - Pt(28), Pt(16), "· " + h, size=8.5, color=SUB)
        y += Pt(16)

    if manual:
        y += Pt(10)
        txt(s, M, y, W - M * 2, Pt(16), "수기 편집 영역", size=10, bold=True, color=BRAND)
        y += Pt(18)
        for m in manual:
            txt(s, M, y, W - M * 2, Pt(16), "· " + m, size=8.5, color=SUB2)
            y += Pt(16)

    foot = []
    if d.get("collectedAt"):
        foot.append("조회일 " + d["collectedAt"])
    if d.get("sources"):
        foot.append("출처: " + " · ".join(d["sources"]))
    if foot:
        txt(s, M, H - Pt(62), W - M * 2, Pt(14), " · ".join(foot), size=8, color=SUB2)
    if d.get("disclaimer"):
        box = card(s, M, H - Pt(46), W - M * 2, Pt(32), fill=CHIP, line=None)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Pt(10)
        tf.margin_top = Pt(5)
        r = tf.paragraphs[0].add_run()
        r.text = d["disclaimer"]
        r.font.size, r.font.color.rgb, r.font.name = Pt(6.5), SUB2, FONT


def build(d, out_path, map_png=None):
    map_png = resolve_map_image(d, map_png)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide_cover(prs, d)
    slide_summary(prs, d, map_png)
    for sec in d.get("sections") or []:
        rows = sec.get("rows") or []
        chunks = [rows[i:i + ROWS_PER_SLIDE] for i in range(0, len(rows), ROWS_PER_SLIDE)] or [[]]
        for i, chunk in enumerate(chunks, start=1):
            slide_section(prs, sec, i, len(chunks), chunk)
    if d.get("wide"):
        slide_wide(prs, d["wide"])
    slide_notes(prs, d)
    prs.save(out_path)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "result.json"
    dst = sys.argv[2] if len(sys.argv) > 2 else "result.pptx"
    png = sys.argv[3] if len(sys.argv) > 3 else None
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    build(data, dst, png)
    print(f"saved: {dst}")
